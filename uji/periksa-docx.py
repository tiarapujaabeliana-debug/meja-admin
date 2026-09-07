"""
Pemeriksaan terakhir: berkas .docx yang kita tulis dibuka oleh pembaca
LAIN, bukan oleh kode kita sendiri.

Ini yang membedakan "susunan ZIP-nya benar menurut kita" dari "Word akan
mau membukanya". Jalankan: python3 uji/periksa-docx.py
"""
import sys, zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

BERKAS = "uji/keluaran/contoh.docx"

def main():
    # 1. Sah sebagai ZIP menurut pustaka standar Python, bukan menurut kita.
    with zipfile.ZipFile(BERKAS) as z:
        rusak = z.testzip()
        assert rusak is None, f"berkas rusak di dalam ZIP: {rusak}"
        print(f"  ok  ZIP sah, {len(z.namelist())} bagian:")
        for n in sorted(z.namelist()):
            print(f"        {n}")

    # 2. Dibuka python-docx — pustaka yang sama sekali tidak tahu kode kita.
    d = Document(BERKAS)
    print(f"\n  ok  python-docx membuka berkasnya, {len(d.paragraphs)} paragraf")
    print(f"  ok  judul dokumen: {d.core_properties.title!r}")
    print(f"  ok  penulis      : {d.core_properties.author!r}\n")

    print("  Isi yang terbaca:")
    for p in d.paragraphs:
        tebal = any(r.bold for r in p.runs if r.bold)
        tengah = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        huruf = {r.font.name for r in p.runs if r.font.name} or {"—"}
        tanda = ("B" if tebal else " ") + ("C" if tengah else " ")
        teks = p.text if p.text.strip() else "(baris kosong)"
        print(f"    [{tanda}] {teks[:58]:<58} {'/'.join(sorted(huruf))}")

    # 3. Hal-hal yang paling gampang rusak, diperiksa satu per satu.
    teks_semua = "\n".join(p.text for p in d.paragraphs)
    cek = [
        ("judul terbaca utuh", "PERJANJIAN KERJA WAKTU TERTENTU" in teks_semua),
        ("nomor surat terbaca", "003/PK/SM/IX/2026" in teks_semua),
        ("NIK terbaca lengkap 16 digit", "3174052809980003" in teks_semua),
        ("spasi perataan tidak dimakan", "Nama       : Rizky Pratama" in teks_semua),
        ("tanda pisah panjang selamat", "—" in teks_semua),
        ("judul dibuat tebal", any(r.bold for p in d.paragraphs for r in p.runs
                                   if "PERJANJIAN" in p.text)),
        ("judul rata tengah", any(p.alignment == WD_ALIGN_PARAGRAPH.CENTER
                                  for p in d.paragraphs if "PERJANJIAN" in p.text)),
        ("kepala pasal TIDAK rata tengah", all(p.alignment != WD_ALIGN_PARAGRAPH.CENTER
                                               for p in d.paragraphs if p.text.startswith("PASAL"))),
        ("baris berspasi rata pakai Courier New", any(r.font.name == "Courier New"
                                                      for p in d.paragraphs for r in p.runs
                                                      if "Rizky" in p.text)),
    ]
    print()
    gagal = 0
    for nama, hasil in cek:
        print(("  ok  " if hasil else "  GAGAL ") + nama)
        if not hasil:
            gagal += 1
    print()
    if gagal:
        print(f"{gagal} pemeriksaan gagal")
        sys.exit(1)
    print("Semua pemeriksaan lolos. Berkasnya sah menurut pembaca di luar kode kita.")

main()
